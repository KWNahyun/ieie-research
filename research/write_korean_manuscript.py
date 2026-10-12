"""Render the full Korean translation, preserving audited result tables."""
import copy
import html
import json
from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parent
OUT = ROOT / 'manuscript'
sys.path.insert(0, str(ROOT / 'vendor/manuscript_deps'))


def main():
    source = json.loads((OUT / 'manuscript_content.json').read_text())
    translated = json.loads((OUT / 'korean_translation.json').read_text())
    doc = copy.deepcopy(source)
    for key in ['title', 'authors', 'abstract', 'keywords']:
        doc[key] = translated[key]
    doc['status'] = '연구 논문 초안 · 저자/소속 및 투고처 미정 · 미투고'
    doc['language'] = 'ko'
    assert len(doc['sections']) == len(translated['sections'])
    table_index = figure_index = 0
    for section, new in zip(doc['sections'], translated['sections']):
        section['title'] = new['title']
        paragraphs = [b for b in section['content'] if b['type'] == 'paragraph']
        assert len(paragraphs) == len(new['paragraphs'])
        for block, text in zip(paragraphs, new['paragraphs']):
            block['text'] = text
        for block in section['content']:
            if block['type'] == 'table':
                block['headers'] = translated['table_headers'][table_index]
                table_index += 1
                block['caption'] = f'표 {table_index}'
                if table_index == 1:
                    for row, label in zip(block['rows'], ['타깃 학습','타깃 검증','타깃 테스트','로컬 소스 전체']):
                        row[0] = label
            elif block['type'] == 'figure':
                block['caption'] = f'그림 {figure_index + 1}. ' + translated['captions'][figure_index]
                figure_index += 1
    assert table_index == 8 and figure_index == 3
    # Verify numerical table cells independently of the translated labels.
    for old, new in zip(source['sections'], doc['sections']):
        for a, b in zip(old['content'], new['content']):
            if a['type'] == 'table':
                for ar, br in zip(a['rows'], b['rows']):
                    assert ar[1:] == br[1:]
            elif a['type'] == 'figure':
                assert a['path'] == b['path'] and (OUT / b['path']).exists()
    (OUT / 'manuscript_ko_content.json').write_text(json.dumps(doc, ensure_ascii=False, indent=2))
    lines = ['# ' + doc['title'], '', doc['authors'], '', doc['status'], '', '## 초록', '', doc['abstract'], '', '주제어: ' + doc['keywords'], '']
    for section in doc['sections']:
        lines += ['## ' + section['title'], '']
        for b in section['content']:
            if b['type'] == 'paragraph':
                lines += [b['text'], '']
            elif b['type'] == 'table':
                lines += [b['caption'], '', '| ' + ' | '.join(b['headers']) + ' |', '| ' + ' | '.join(['---'] * len(b['headers'])) + ' |']
                lines += ['| ' + ' | '.join(r) + ' |' for r in b['rows']]
                lines += ['']
            else:
                lines += [f"![{b['caption']}]({b['path']})", '', b['caption'], '']
    lines += ['## 참고문헌', '']
    for i, r in enumerate(doc['references'], 1):
        lines += [reference(i, r), '']
    (OUT / 'manuscript_ko.md').write_text('\n'.join(lines))
    word(doc)
    pdf(doc)
    tex(doc)
    (OUT / 'README_ko.md').write_text('한글 논문 초안입니다. 영문 초안의 모든 절을 번역하고 표 8개, 그림 3개 및 참고문헌 8개를 유지했습니다.\n\n수정용: manuscript_ko.md / manuscript_ko.docx\n검토용 PDF: manuscript_ko_review.pdf\nLaTeX: manuscript_ko.tex (XeLaTeX와 kotex 필요, 이 환경에서는 컴파일하지 않음)\n\n저자·소속 및 투고처는 미정입니다. PDF와 Word는 검토용으로, 특정 학회의 최종 투고 양식을 적용한 파일은 아닙니다. 그림 내부의 영문 레이블은 원본을 유지했습니다. 수치 표는 검증된 영문 결과에서 복사했습니다. 추가 비교의 탐색적 성격과 물리 단위·지리적 독립성의 미확인 사항을 명시했습니다.\n')
    (OUT / 'korean_translation_verification.json').write_text(json.dumps({'status':'PASS','sections':len(doc['sections']),'paragraphs':sum(len(s['paragraphs']) for s in translated['sections']), 'tables':table_index,'figures':figure_index,'references':len(doc['references']),'numerical_table_cells':'identical to English source','figure_paths':'identical to English source'}, indent=2))
    print('Korean manuscript written: Markdown, DOCX, PDF, LaTeX source; tables verified')


def reference(i, r):
    return f"[{i}] {r['authors']}. {r['title']}. {r.get('venue','arXiv')}, {r['year']}. {r['url']}"


def word(d):
    from docx import Document
    from docx.shared import Inches, Pt
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    document = Document()
    for name in ['Normal','Title','Heading 1','Heading 2','Caption']:
        style = document.styles[name]
        style.font.name = 'Noto Sans CJK KR'
        fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        fonts.set(qn('w:eastAsia'), 'Noto Sans CJK KR')
    document.styles['Normal'].font.size = Pt(10)
    document.add_heading(d['title'], 0)
    document.add_paragraph(d['authors'])
    document.add_paragraph(d['status'])
    document.add_heading(d.get('abstract_label', '초록'), 1)
    document.add_paragraph(d['abstract'])
    document.add_paragraph('주제어: ' + d['keywords'])
    for s in d['sections']:
        document.add_heading(s['title'], 1)
        for b in s['content']:
            if b['type'] == 'paragraph':
                document.add_paragraph(b['text'])
            elif b['type'] == 'table':
                document.add_paragraph(b['caption'], style='Caption')
                t = document.add_table(rows=1, cols=len(b['headers']))
                t.style = 'Table Grid'
                for cell, text in zip(t.rows[0].cells, b['headers']):
                    cell.text = text
                repeat = OxmlElement('w:tblHeader')
                t.rows[0]._tr.get_or_add_trPr().append(repeat)
                for row in b['rows']:
                    for cell, text in zip(t.add_row().cells, row):
                        cell.text = text
            else:
                document.add_picture(str(OUT / b['path']), width=Inches(6.1))
                document.add_paragraph(b['caption'], style='Caption')
    document.add_heading('참고문헌', 1)
    for i, r in enumerate(d['references'], 1):
        document.add_paragraph(reference(i, r))
    document.save(OUT / 'manuscript_ko.docx')


def pdf(d):
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib import colors
    from reportlab.lib.utils import ImageReader
    # Embed a TrueType subset converted from the installed Noto CFF collection.
    # This avoids viewer-dependent Korean CID substitutions and missing symbols.
    from fontTools.ttLib import TTFont as SourceFont
    from fontTools.fontBuilder import FontBuilder
    from fontTools.pens.ttGlyphPen import TTGlyphPen
    from fontTools.pens.cu2quPen import Cu2QuPen
    installed = SourceFont('/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc', fontNumber=1)
    cmap = installed.getBestCmap()
    chars = set(json.dumps(d, ensure_ascii=False) + '한글 연구 논문 초안 · 정규화 영상 평가0123456789')
    codepoints = {ord(c): cmap[ord(c)] for c in chars if ord(c) in cmap}
    order = ['.notdef'] + sorted(set(codepoints.values()) - {'.notdef'})
    glyph_set = installed.getGlyphSet()
    glyphs = {}
    for name in order:
        pen = TTGlyphPen(None)
        glyph_set[name].draw(Cu2QuPen(pen, max_err=1.0, reverse_direction=True))
        glyphs[name] = pen.glyph()
    builder = FontBuilder(installed['head'].unitsPerEm, isTTF=True)
    builder.setupGlyphOrder(order)
    builder.setupCharacterMap(codepoints)
    builder.setupGlyf(glyphs)
    builder.setupHorizontalMetrics({name: installed['hmtx'].metrics[name] for name in order})
    builder.setupHorizontalHeader(ascent=installed['hhea'].ascent, descent=installed['hhea'].descent)
    builder.setupNameTable({'familyName':'IEIE Noto Korean Subset','styleName':'Regular','uniqueFontIdentifier':'IEIE-Noto-KO','fullName':'IEIE Noto Korean Subset','psName':'IEIENotoKO'})
    builder.setupOS2(sTypoAscender=installed['hhea'].ascent, sTypoDescender=installed['hhea'].descent, usWinAscent=installed['hhea'].ascent, usWinDescent=abs(installed['hhea'].descent))
    builder.setupPost()
    font_path = Path('/tmp/ieie_noto_korean_subset.ttf')
    builder.save(font_path)
    font = 'EmbeddedNotoKO'
    pdfmetrics.registerFont(TTFont(font, str(font_path)))
    body = ParagraphStyle('BodyKO', fontName=font, fontSize=10, leading=16, wordWrap='CJK', spaceAfter=9)
    heading = ParagraphStyle('HeadingKO', parent=body, fontSize=13, leading=19, spaceBefore=12, spaceAfter=9)
    title = ParagraphStyle('TitleKO', parent=body, fontSize=18, leading=27, spaceAfter=16)
    caption = ParagraphStyle('CaptionKO', parent=body, fontSize=8, leading=12, spaceAfter=8)
    cell = ParagraphStyle('CellKO', parent=caption, fontSize=7.5, leading=11, spaceAfter=0)
    p = lambda text, style=body: Paragraph(html.escape(text), style)
    story = [p(d['title'], title), p(d['authors']), p(d['status'], caption), p(d.get('abstract_label', '초록'), heading), p(d['abstract']), p('주제어: ' + d['keywords'])]
    for s in d['sections']:
        story.append(p(s['title'], heading))
        for b in s['content']:
            if b['type'] == 'paragraph':
                story.append(p(b['text']))
            elif b['type'] == 'table':
                story.append(p(b['caption'], caption))
                t = Table([[p(x, cell) for x in row] for row in [b['headers']] + b['rows']], colWidths=[499 / len(b['headers'])] * len(b['headers']), repeatRows=1)
                t.setStyle(TableStyle([('GRID',(0,0),(-1,-1),.3,colors.grey),('BACKGROUND',(0,0),(-1,0),colors.HexColor('#eeeeee')),('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),4),('RIGHTPADDING',(0,0),(-1,-1),4)]))
                story += [t, Spacer(1, 12)]
            else:
                w, h = ImageReader(str(OUT / b['path'])).getSize()
                story += [Image(str(OUT / b['path']), width=499, height=499*h/w), p(b['caption'], caption)]
    story.append(p('참고문헌', heading))
    for i, r in enumerate(d['references'], 1):
        story.append(p(reference(i, r), caption))
    def page(canvas, doc):
        canvas.setFont(font, 8)
        canvas.drawString(48, 24, d.get('footer', '한글 연구 논문 초안 · 정규화 영상 평가'))
        canvas.drawRightString(547, 24, str(doc.page))
    document = SimpleDocTemplate(str(OUT / 'manuscript_ko_review.pdf'), pagesize=(595.28,841.89), leftMargin=48, rightMargin=48, topMargin=45, bottomMargin=45, title=d['title'], author=d['authors'])
    document.build(story, onFirstPage=page, onLaterPages=page)


def tex(d):
    def esc(text):
        text = text.replace('⁻⁴', r'$^{-4}$').replace('⊙', r'$\odot$')
        for a,b in [('&',r'\&'),('%',r'\%'),('_',r'\_'),('#',r'\#')]:
            text = text.replace(a,b)
        return text
    lines = [r'\documentclass[10pt,a4paper]{article}',r'\usepackage{kotex,fontspec,graphicx,booktabs,amsmath,url,geometry}',r'\geometry{margin=22mm}',r'\setmainhangulfont{Noto Serif CJK KR}',r'\setsanshangulfont{Noto Sans CJK KR}',r'\title{'+esc(d['title'])+'}',r'\author{'+esc(d['authors'])+'}',r'\date{}',r'\begin{document}',r'\maketitle',r'\begin{abstract}',esc(d['abstract']),r'\end{abstract}',r'\noindent 주제어: '+esc(d['keywords'])]
    for s in d['sections']:
        lines.append(r'\section*{'+esc(s['title'])+'}')
        for b in s['content']:
            if b['type'] == 'paragraph':
                lines += [esc(b['text']), '']
            elif b['type'] == 'table':
                lines += [r'\begin{table}[htbp]\centering\small', r'\caption{'+esc(b['caption'])+'}',r'\resizebox{\textwidth}{!}{\begin{tabular}{'+'l'*len(b['headers'])+'}',r'\toprule',' & '.join(esc(x) for x in b['headers'])+r' \\',r'\midrule']
                lines += [' & '.join(esc(x) for x in row)+r' \\' for row in b['rows']]
                lines += [r'\bottomrule\end{tabular}}\end{table}']
            else:
                lines += [r'\begin{figure}[htbp]\centering',r'\includegraphics[width=\textwidth]{'+b['path']+'}',r'\caption{'+esc(b['caption'].split('. ',1)[-1])+'}',r'\end{figure}']
    lines += [r'\begin{thebibliography}{99}']
    for r in d['references']:
        lines += [r'\bibitem{'+r['key']+'} '+esc(f"{r['authors']}. {r['title']}. {r.get('venue','arXiv')}, {r['year']}.")+r' \url{'+r['url']+'}']
    lines += [r'\end{thebibliography}',r'\end{document}']
    (OUT / 'manuscript_ko.tex').write_text('\n'.join(lines)+'\n')


if __name__ == '__main__':
    main()
