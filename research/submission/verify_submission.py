"""Check rendered length, preserved native equations/images and numeric tables."""
import sys,json,hashlib,subprocess,re
from pathlib import Path
from lxml import etree
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'vendor/manuscript_deps'))
from docx import Document
P=Path(__file__).resolve().parent;ROOT=P.parents[1]
team=Document(ROOT/'DINOv2_본론_구현_초안.docx');merged=Document(P/'DINOv2_CNN_본론_통합검토본.docx');own=Document(P/'CNN_담당본론_초안.docx')
def maths(doc):return [etree.tostring(x,method='c14n') for x in doc._element.xpath('.//m:oMath')]
def images(doc):return {hashlib.sha256(r.target_part.blob).hexdigest() for r in doc.part.rels.values() if r.reltype.endswith('/image')}
assert maths(team)==maths(merged)
assert images(team)<=images(merged)
assert len(merged.inline_shapes)==4
assert [p.text for p in merged.paragraphs if p.text in ['2. 본론','3. 구현']]==['2. 본론','3. 구현']
assert all(p in '\n'.join(x.text for x in merged.paragraphs) for p in ['2.4 CNN 기반 조건부 인코더 전이','3.3 CNN 인코더 전이 성능 비교','3.4 실험 결과 종합 및 논의'])
pages={}
for stem,wanted in [('CNN_담당본론_초안',2),('DINOv2_CNN_본론_통합검토본',4)]:
 info=subprocess.check_output(['pdfinfo',str(P/(stem+'.pdf'))],text=True)
 pages[stem]=int(re.search(r'Pages:\s+(\d+)',info).group(1));assert pages[stem]==wanted
 text=subprocess.check_output(['pdftotext',str(P/(stem+'.pdf')),'-'],text=True)
 assert '평가 대기' not in text and '기입 위치' not in text
for mode in ['base','encoder']:
 metrics=json.loads((P.parent/f'runs/cnn_matched_v1/{mode}_seed42/test_metrics.json').read_text())['metrics']
 for d in [own,merged]:assert any(f"{metrics['mae']:.5f}" in c.text for t in d.tables for r in t.rows for c in r.cells)

checks={'status':'PASS','pdf_pages':pages,'preserved_native_equations':len(maths(team)),'preserved_teammate_images':len(images(team)),'merged_images':len(merged.inline_shapes),'numeric_tables_match_analysis':True,'no_pending_result_placeholders':True,'original_teammate_sha256':hashlib.sha256((ROOT/'DINOv2_본론_구현_초안.docx').read_bytes()).hexdigest(),'template_sha256':hashlib.sha256((ROOT/'논문양식.docx').read_bytes()).hexdigest()}
(P/'artifact_verification.json').write_text(json.dumps(checks,ensure_ascii=False,indent=2));print(json.dumps(checks,ensure_ascii=False,indent=2))
