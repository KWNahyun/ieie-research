"""Apply evidence-bounded content, notation and reference revisions."""
import copy,re
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

def edit(d):
 def find(prefix):return next(p for p in d.paragraphs if p.text.strip().startswith(prefix))
 def replace(p,text):
  assert not p._p.xpath('.//m:oMath | .//w:drawing')
  p.runs[0].text=text
  for r in p.runs[1:]:r.text=''
 def remove(p):p._p.getparent().remove(p._p)
 def insert_after(p,text):
  node=OxmlElement('w:p')
  if p._p.pPr is not None:node.append(copy.deepcopy(p._p.pPr))
  p._p.addnext(node);new=Paragraph(node,p._parent);r=new.add_run(text)
  if p.runs and p.runs[0]._r.rPr is not None:r._r.insert(0,copy.deepcopy(p.runs[0]._r.rPr))
  return new
 abstract=find('Directional electromagnetic field')
 replace(abstract,'The directional electromagnetic field (EMF) dataset contains limited training samples, motivating additional pretraining on RadioMapSeer. Domain differences, however, make the transferability of learned representations uncertain. We investigate transfer learning using the pretrained image foundation model DINOv2 and a randomly initialized convolutional neural network (CNN). The representative DINOv2 configuration is selected by the lowest validation MAE among the compared configurations. Encoder–decoder transfer with low-rank adaptation and L1 loss in both domains reduces test MAE by 8.0% relative to DINOv2 without additional RadioMapSeer pretraining. In the CNN comparison, source encoder initialization increases test MAE by 3.9%, while SSIM remains similar. Encoder-only transfer also increases MAE in the DINOv2 L1-to-L1 setting. These observations suggest that transfer benefits depend on the transfer scope and adaptation strategy in the evaluated setting.')
 intro=find('지향성 전자파(EMF) 지도 예측에는')
 replace(intro,intro.text.replace('RadioMapSeer[1]','RadioMapSeer[1][11]').replace('사용하는 걸','사용하는 것을'))
 remove(find('본 연구는 RadioMapSeer[1] 소스 과제에서'))
 cnn=find('일반 영상 사전학습 없이 시작하는 모델에서 소스 인코더')
 replace(cnn,'일반 영상 사전학습 없이 시작하는 모델에서 소스 인코더 초기화의 효과를 확인하기 위한 보조 비교에는 CNN 기반 예측기를 사용했다.')
 # Initialization control belongs in the transfer section; seeds only in setup.
 paired=find('Source와 Target 시드는 각각 142·42로 설정했다.')
 replace(paired,'BASE와 ENC의 디코더, 출력층, 패턴 투영층, FiLM은 동일하게 초기화했다. ENC는 소스 인코더만 복사한 뒤 전체 모델을 학습했다.')
 data=find('RadioMapSeer DPM/no-cars의 building map 700')
 padding=find('DINOv2 기반 모델의 Source와 Target 입력은')
 replace(padding,'DINOv2 기반 모델의 소스와 타깃 입력은 replication padding으로 각각 256×256에서 266×266, 500×500에서 518×518로 확장했다. 예측 결과는 원래 크기로 잘랐으며, 정답은 0~1 범위로 정규화했다.')
 shared=find('검증 MAE가 가장 낮은 모델을 선택해 평가용 데이터에서')
 replace(shared,'두 모델 모두 검증 MAE가 가장 낮은 체크포인트를 선택해 테스트 데이터에서 평가했다. MAE·RMSE·PSNR은 전체 유효 픽셀로 계산하고, SSIM[9]은 영상별 값을 평균했다.')
 extra=find('CNN 기반 모델도 AdamW와 cosine schedule을')
 replace(extra,'CNN 기반 모델도 동일한 최적화·학습률 일정·최대 epoch·조기 종료 설정을 사용했다. 증강·EMA·출력 clipping은 적용하지 않았다.')
 first_results=find('Source 사전학습에서는 MAE')
 replace(first_results,'DINOv2의 소스 사전학습에서는 MAE·RMSE·PSNR은 L1, SSIM은 M1이 더 좋았다(표 2).')
 desc=find('표 3에서 L1 → L1은')
 replace(desc,'표 3에서 L1→L1은 소스와 타깃에 모두 L1 손실을 사용한 구성이다. 대표 구성은 비교한 타깃 구성 중 검증 MAE가 가장 낮은 인코더·디코더 전이+LoRA의 L1→L1로 선택했다(반올림 전 검증 MAE 0.05918). 이 구성은 추가 소스 학습이 없는 타깃 L1 대비 검증·테스트 MAE를 각각 4.4%, 8.0% 낮췄고 RMSE·PSNR도 개선했다.')
 results=find('평가 MAE와 SSIM은 인코더·디코더 전이의')
 replace(results,'테스트 MAE와 SSIM은 인코더·디코더 전이의 L1→M1에서 가장 좋았으나, 대표 구성은 테스트 지표가 아닌 검증 MAE를 기준으로 정했다. 인코더 단독 전이는 추가 소스 학습이 없는 동일 타깃 손실 조건보다 검증 MAE가 모두 증가했고, 테스트 MAE는 L1→L1에서만 증가했다. 인코더·디코더 전이+LoRA는 동일 손실 조합의 인코더 단독 전이보다 모든 보고 지표가 개선됐으며, 검증·테스트 MAE는 각각 2.8~6.6%, 2.3~10.5% 감소했다.')
 cnn_results=find('동일한 Target 시드 42에서 CNN의')
 replace(cnn_results,'동일한 타깃 시드 42에서 CNN의 BASE와 ENC의 테스트 MAE는 각각 0.0796, 0.0828이었다. 소스 인코더 초기화 후 MAE는 3.9% 증가했다. RMSE와 PSNR은 BASE가 좋았고, SSIM은 0.8998과 0.9014로 비슷했다. 이 조건에서는 소스 인코더 전이에 따른 픽셀별 예측 오차 감소가 관찰되지 않았다.')
 conclusion=find('본 논문은 지향성 안테나의 EMF map prediction')
 replace(conclusion,'본 논문은 지향성 EMF 영상 예측에서 RadioMapSeer 추가 사전학습의 효과를 비교하였다. 검증 MAE를 기준으로 선택한 인코더·디코더 전이+LoRA의 L1→L1 구성은 추가 사전학습이 없는 구성보다 테스트 MAE가 8.0% 낮았다. 데이터셋의 도메인 차이가 있음에도 이 조건에서 개선이 관찰되었다.')
 tail=find('한편 CNN의 보조 비교에서는')
 replace(tail,'CNN의 보조 비교에서는 소스 인코더 초기화 후 테스트 MAE가 3.9% 증가했고 SSIM은 비슷했다. DINOv2의 L1→L1 인코더 단독 전이에서도 MAE가 증가했다. 결과는 소스 인코더 초기화만으로 개선이 보장되지 않으며, 관찰된 전이 효과가 전이 범위와 적응 방식에 따라 달라짐을 시사한다. 해석 범위는 제공된 정규화 영상과 비교한 학습 조건으로 한정한다.')
 spatial=find('S는 building map')
 replace(spatial,'S는 건물 점유 영상, 송신기 Gaussian heatmap(σ=10), log-distance map으로 구성된다. 점유값은 소스 건물 PNG가 0이 아닌 픽셀, 타깃 건물 영상의 B채널이 0인 픽셀을 1로 정의했다.')
 loss=find('손실은 건물 내부와 결측 위치를 제외한')
 for r in loss.runs:
  if r._element.xpath('.//m:oMath'):continue
  r.text=r.text.replace('손실은 건물 내부와 결측 위치를 제외한 유효 픽셀에서 계산했다.','손실은 건물 영상의 점유 마스크를 제외한 픽셀에서 계산했다.')
 # Define equation (1)'s angles and fix equation (2) as native OMML structures.
 wrap=find('wrap은 상대각을')
 replace(wrap,'a는 안테나 방위각(azimuth), b는 빔폭(beamwidth)이며 단위는 도(°)이다. wrap은 상대각을 [−180°,180°)로 변환한다. 송신기와 일치하는 픽셀의 n2i 값은 1이다.')
 equation=next(m for m in d._element.xpath('.//m:oMath') if any('F̃' in t.text for t in m.findall('.//'+qn('m:t')) if t.text))
 for el in list(equation):equation.remove(el)
 def run(s):
  r=OxmlElement('m:r');pr=OxmlElement('w:rPr');f=OxmlElement('w:rFonts');f.set(qn('w:ascii'),'Cambria Math');f.set(qn('w:hAnsi'),'Cambria Math');pr.append(f);sz=OxmlElement('w:sz');sz.set(qn('w:val'),'18');pr.append(sz);r.append(pr);t=OxmlElement('m:t');t.text=s;t.set(qn('xml:space'),'preserve');r.append(t);return r
 def sub(base):
  s=OxmlElement('m:sSub');e=OxmlElement('m:e');e.append(base);s.append(e);i=OxmlElement('m:sub');i.append(run('i'));s.append(i);return s
 acc=OxmlElement('m:acc');pr=OxmlElement('m:accPr');char=OxmlElement('m:chr');char.set(qn('m:val'),'̃');pr.append(char);acc.append(pr);e=OxmlElement('m:e');e.append(run('F'));acc.append(e)
 for node in [sub(acc),run(' = (1 + '),sub(run('γ')),run('(Q)) ⊙ '),sub(run('F')),run(' + '),sub(run('β')),run('(Q)     (2)')]:equation.append(node)
 # Use a native absolute-value delimiter in equation (1), avoiding renderer parsing of a text bar.
 eq1=d._element.xpath('.//m:oMath')[0]
 condition=next(e for e in eq1.findall('.//'+qn('m:e')) if any('|wrap' in (t.text or '') for t in e.findall('./'+qn('m:r')+'/'+qn('m:t'))))
 old=next(r for r in condition.findall(qn('m:r')) if any('|wrap' in (t.text or '') for t in r.findall(qn('m:t'))))
 absolute=OxmlElement('m:d');pr=OxmlElement('m:dPr')
 for tag in ['begChr','endChr']:
  el=OxmlElement('m:'+tag);el.set(qn('m:val'),'|');pr.append(el)
 absolute.append(pr);e=OxmlElement('m:e');e.append(run('wrap(θ(u) − a)'));absolute.append(e)
 idx=list(condition).index(old);condition.remove(old);condition.insert(idx,absolute);condition.insert(idx+1,run(' ≤ '))
 # Add RadioUNet, then renumber every citation by first occurrence.
 references=[p for p in d.paragraphs if re.match(r'^\[\d+\]',p.text)]
 unet=next(p for p in references if p.text.startswith('[10]'))
 replace(unet,'[10] O. Ronneberger, P. Fischer, and T. Brox, “U-Net: Convolutional Networks for Biomedical Image Segmentation,” in Proc. MICCAI, LNCS, vol. 9351, pp. 234–241, 2015.')
 radio=insert_after(unet,'[11] R. Levie, Ç. Yapar, G. Kutyniok, and G. Caire, “RadioUNet: Fast Radio Map Estimation With Convolutional Neural Networks,” IEEE Transactions on Wireless Communications, vol. 20, no. 6, pp. 4001–4015, Jun. 2021.')
 reference_nodes={int(re.match(r'^\[(\d+)\]',p.text)[1]):p._p for p in references+[radio]}
 bibliography=find('참고문헌')._p
 mapping={}
 for node in d._element.body:
  if node is bibliography:break
  for t in node.xpath('.//w:t'):
   for match in re.finditer(r'\[(\d+)\]',t.text or ''):
    num=int(match[1])
    if num not in mapping:mapping[num]=len(mapping)+1
 assert set(mapping)==set(reference_nodes),(mapping,reference_nodes.keys())
 for t in d._element.xpath('.//w:t'):
  if t.text:t.text=re.sub(r'\[(\d+)\]',lambda m:'['+str(mapping[int(m[1])])+']',t.text)
 sect=d._element.body.find(qn('w:sectPr'))
 for old in sorted(mapping,key=mapping.get):sect.addprevious(reference_nodes[old])
 # Normalize terminology in Korean prose/tables; keep English abstract and bibliography.
 in_abstract=False;in_refs=False
 for node in d._element.body:
  text=''.join(node.xpath('.//w:t/text()'))
  if text=='Abstract':in_abstract=True
  if text=='1. 서론':in_abstract=False
  if text=='참고문헌':in_refs=True
  if in_abstract or in_refs:continue
  for t in node.xpath('.//w:t'):
   if not t.text:continue
   s=t.text.replace('사후 학습 시','타깃 학습 시').replace('Source','소스').replace('Target','타깃').replace('test MAE','테스트 MAE').replace('test set','테스트 데이터').replace('train set','학습 데이터')
   s=s.replace('소스(소스)','소스(Source)').replace('타깃(타깃)','타깃(Target)').replace('3. 구현','3. 실험')
   s=re.sub(r'(\d+\.\d+)%',lambda m:f'{float(m[1]):.1f}%',s)
   t.text=s
 # All metric decimals: 4 places; PSNR retains 2 places. Raw files unchanged.
 for table in d.tables[3:]:
  for row in table.rows[1:]:
   for ci,c in enumerate(row.cells):
    if re.fullmatch(r'\d+\.\d+',c.text):
     p=c.paragraphs[0];r=p.runs[0];r.text=f'{float(c.text):.{2 if ci==len(row.cells)-1 else 4}f}'
     for other in p.runs[1:]:other.text=''
 return mapping
