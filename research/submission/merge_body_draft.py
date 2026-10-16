"""Build a review copy, preserving the teammate's equations/images and original file."""
import sys,copy,io
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'vendor/manuscript_deps'))
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt,Inches
ROOT=Path(__file__).resolve().parents[2];OUT=Path(__file__).resolve().parent
doc=Document(ROOT/'DINOv2_본론_구현_초안.docx');own=Document(OUT/'CNN_담당본론_초안.docx')
def text(e):return ''.join(e.itertext()) if False else ''.join(e.xpath('.//w:t/text()'))
def find(s):return next(p._p for p in doc.paragraphs if p.text==s)
anchor_method=find('Ⅲ. 구현');anchor_result=find('참고문헌')
for p in doc.paragraphs:
 if p.text.startswith('RadioMapSeer[1](Source)'):
  for r in p.runs:r.text=''
  p.add_run('본 연구는 RadioMapSeer[1] 소스 과제에서 학습한 지식이 지향성 EMF 지도 예측에 기여하는 조건을 검토한다. DINOv2[2]에서는 전이 범위와 손실 조합을 비교하고, CNN에서는 동일한 입력 정보와 데이터 분할을 이용해 인코더 전이의 효과를 비교한다.')
 if p.text=='2.2 모델 아키텍처':
  p.runs[0].text='2.2 DINOv2 기반 예측 모델'
 if p.text=='3.2 사전학습 및 전이 성능 비교':
  p.runs[0].text='3.2 DINOv2 사전학습 및 전이 성능 비교'
 if p.text.startswith('Source와 Target 입력은 replication padding'):
  p.runs[0].text='DINOv2 실험에서 '+p.runs[0].text
 if p.text.startswith('Source와 Target은 배치 크기'):
  p.runs[0].text='DINOv2 실험에서 '+p.runs[0].text
 for r in p.runs:
  if '디코더를 새로 초기화해 전체 학습했다.' in r.text:r.text=r.text.replace('디코더를 새로 초기화해 전체 학습했다.','디코더를 새로 초기화해 학습했다.')

def add(e,anchor):
 n=copy.deepcopy(e)
 for el in n.iter():
  if el.tag==qn('w:sectPr'):
   el.getparent().remove(el)
  for attr in [qn('r:embed'),qn('r:link')]:
   old=el.get(attr)
   if old:
    rel=own.part.rels[old]
    if rel.reltype.endswith('/image'):
     rid,_=doc.part.get_or_add_image(io.BytesIO(rel.target_part.blob));el.set(attr,rid)
  if el.tag.endswith('}docPr'):el.set('id','4')
  if el.tag==qn('w:t'):
   el.text=(el.text or '').replace('[1]','[10]').replace('[2]','[5]').replace('그림 1.','그림 4.').replace('표 2.','표 4.')
 for r in n.xpath('.//w:r'):
  pr=r.find(qn('w:rPr'))
  if pr is None:
   from docx.oxml import OxmlElement
   pr=OxmlElement('w:rPr');r.insert(0,pr)
  fonts=pr.find(qn('w:rFonts'))
  if fonts is None:
   from docx.oxml import OxmlElement
   fonts=OxmlElement('w:rFonts');pr.insert(0,fonts)
  fonts.set(qn('w:eastAsia'),'HY신명조');fonts.set(qn('w:ascii'),'HY신명조')
 anchor.addprevious(n)
 return n

def p(s,anchor,heading=False):
 x=doc.add_paragraph();x.paragraph_format.line_spacing=1.2;x.paragraph_format.space_after=Pt(3)
 x.paragraph_format.keep_with_next=heading
 r=x.add_run(s);r.font.size=Pt(10 if heading else 9);r.bold=heading
 r.font.name='HY신명조';r._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'HY신명조')
 anchor.addprevious(x._p)
 return x

p('2.4 CNN 기반 조건부 인코더 전이',anchor_method,True)
p('DINOv2의 전이 범위 비교에 추가하여 CNN에서도 인코더 전이의 효과가 나타나는지 검토했다. CNN은 2.1절과 동일한 S·P·Q를 입력받는다. 동일 구조에서 무작위 초기화와 Source 인코더 초기화를 비교했으며, 디코더 전이와 손실 조합은 DINOv2 실험에서 다뤘다.',anchor_method)
selected=False
for e in own._element.body:
 t=text(e)
 if t=='3.2 조건부 CNN 구조':selected=True;continue
 if e.xpath('.//w:br[@w:type="page"]'):break
 if selected:
  if t.startswith('F′ ='):continue
  n=add(e,anchor_method)
  for el in n.xpath('.//w:t'):
   if el.text=='3.3 인코더 전이와 대응 초기화':el.text='대응 초기화 및 소스 전이'
   if el.text:el.text=el.text.replace('조건화는 다음과 같이 계산한다.','조건화는 식 (2)와 같이 계산한다.')

p('3.3 CNN 인코더 전이 성능 비교',anchor_result,True)
p('CNN은 3.1절과 동일한 데이터 분할과 평가 함수를 사용했다. Source와 Target 시드는 각각 142·42로 설정했다. CNN의 BASE는 무작위 초기화 모델이며, 모델별 설정 차이는 표 3과 같다.',anchor_result)
p('표 3. 모델별 실험 구성 차이',anchor_result,True)
t=doc.add_table(rows=1,cols=3);t.autofit=False
for col in t.columns:col.width=Inches(1.05)
for c,v in zip(t.rows[0].cells,['항목','DINOv2','CNN']):c.text=v
rows=[['타깃 기준 초기화','공개 사전학습','무작위'],['Source 배치','32','8'],['Target 배치','8','8'],['주 비교 Target 시드','42','42'],['Source 입력 크기','266×266','256×256'],['Target 입력 크기','518×518','504×504'],['패딩','replication','상수 0'],['Target 인코더 적응','LoRA 등','전체 미세조정'],['Source→Target 손실','L1/M1 조합','L1→L1']]
for row in rows:
 for c,v in zip(t.add_row().cells,row):c.text=v
for row in t.rows:
 for c in row.cells:
  c.width=Inches(1.05)
  for x in c.paragraphs:
   x.paragraph_format.space_after=Pt(2)
   for r in x.runs:r.font.size=Pt(8)
anchor_result.addprevious(t._tbl)
p('최적화는 AdamW와 cosine schedule을 사용했다. 학습률, warmup, 최대 epoch, 조기 종료 대기 기간은 3.1절과 동일하게 설정했다. 증강·EMA·출력 clipping은 적용하지 않았다. Validation MAE가 가장 낮은 모델을 선택하여 test set에서 성능을 측정했다.',anchor_result)
selected=False
for e in own._element.body:
 t=text(e)
 if t=='4.2 평가 지표와 반복 실험 결과':selected=True;continue
 if t=='참고문헌':break
 if selected:
  n=add(e,anchor_result)
  for el in n.xpath('.//w:t'):
   if el.text=='4.3 DINOv2 비교의 해석 범위':el.text='3.4 실험 결과 종합 및 논의'
# Append U-Net after the teammate's existing nine references.
p('[10] O. Ronneberger, P. Fischer, and T. Brox, “U-Net: Convolutional Networks for Biomedical Image Segmentation,” arXiv:1505.04597, 2015.',doc._element.body.find(qn('w:sectPr')))
# Remove manual pagination and use the conference template's body leading.
for para in list(doc.paragraphs):
 for br in list(para._p.xpath('.//w:br[@w:type="page" or @w:type="column"]')):br.getparent().remove(br)
 para.paragraph_format.page_break_before=False
 para.paragraph_format.line_spacing=1.2
 if not para.text.strip() and not para._p.xpath('.//w:drawing | .//w:pict | .//m:oMath | .//m:oMathPara | .//w:sectPr'):
  para._p.getparent().remove(para._p)
for para in doc.paragraphs:
 if para.text=='II. 본론':para.runs[0].text='2. 본론'
 if para.text=='Ⅲ. 구현':para.runs[0].text='3. 구현'
 for run in para.runs:
  if run._element.xpath('.//w:drawing | .//w:pict | .//m:oMath | .//w:object'):continue
  run.text=run.text.replace('인코더·디코더 전이는', '인코더·디코더 전이+LoRA 조합은') if run.text else run.text
# Do not change equations or drawing runs while normalizing text.
doc.core_properties.title='DINOv2·CNN 본론 통합 검토본';doc.core_properties.author=''
path=OUT/'DINOv2_CNN_본론_통합검토본.docx';doc.save(path);print(path)
