"""Merge supplied front matter and conclusion with the verified DINO/CNN body."""
import copy, io, sys, hashlib, json
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'vendor/manuscript_deps'))
from docx import Document
from docx.oxml.ns import qn
from lxml import etree
OUT=Path(__file__).resolve().parent
ROOT=OUT.parents[1]
front_path=ROOT/'DINOv2_초록_서론_결론.docx'
body_path=OUT/'DINOv2_CNN_본론_통합검토본.docx'
front=Document(front_path); body=Document(body_path)
original_hashes={str(p.relative_to(ROOT)):hashlib.sha256(p.read_bytes()).hexdigest() for p in [front_path,body_path]}
conclusion=next(p._p for p in front.paragraphs if p.text=='Ⅳ. 결론 및 향후 연구 방향')
# The source's section break before the conclusion becomes the break before the body.
ref_nodes=[]
is_ref=False
for node in body._element.body:
 if node.tag==qn('w:sectPr'):continue
 txt=''.join(node.xpath('.//w:t/text()'))
 if txt=='참고문헌':is_ref=True
 new=copy.deepcopy(node)
 for el in new.iter():
  for attr in [qn('r:embed'),qn('r:link'),qn('r:id')]:
   rid=el.get(attr)
   if not rid:continue
   rel=body.part.rels[rid]
   if rel.reltype.endswith('/image'):
    newrid,_=front.part.get_or_add_image(io.BytesIO(rel.target_part.blob))
   elif rel.is_external:
    newrid=front.part.relate_to(rel.target_ref,rel.reltype,is_external=True)
   else:raise ValueError(f'Unsupported relationship: {rel.reltype}')
   el.set(attr,newrid)
 if txt=='2.2 DINOv2 기반 예측 모델2.2 모델 아키텍처':
  texts=new.xpath('.//w:t');texts[0].text='2.2 DINOv2 기반 예측 모델'
  for t in texts[1:]:t.text=''
 if is_ref:ref_nodes.append(new)
 else:conclusion.addprevious(new)
for p in front.paragraphs:
 if p.text=='I. 서론':
  for i,r in enumerate(p.runs):r.text='1. 서론' if i==0 else ''
 if p.text=='Ⅳ. 결론 및 향후 연구 방향':
  for i,r in enumerate(p.runs):r.text='4. 결론 및 향후 연구 방향' if i==0 else ''
sect=front._element.body.find(qn('w:sectPr'))
for n in ref_nodes:sect.addprevious(n)
for i,el in enumerate(front._element.xpath('.//wp:docPr'),1):el.set('id',str(i))
front.core_properties.title='지향성 전자파 지도 예측을 위한 RadioMapSeer 기반 교차 도메인 전이학습'
from refine_final_draft import refine
refine(front)
from reviewer_revision import edit
citation_mapping=edit(front)
final=ROOT/'DINOv2_CNN_최종초안.docx';front.save(final)
check=Document(final)
def maths(d):return [etree.tostring(x,method='c14n') for x in d._element.xpath('.//m:oMath')]
def imgs(d):return {hashlib.sha256(r.target_part.blob).hexdigest() for r in d.part.rels.values() if r.reltype.endswith('/image')}
assert len(maths(check))==len(maths(body))==7
assert [m for i,m in enumerate(maths(check)) if i not in [0,1]]==[m for i,m in enumerate(maths(body)) if i not in [0,1]]
assert len(check._element.xpath('.//m:acc'))==1
assert len(check._element.xpath('.//m:oMath')[1].findall('.//'+qn('m:sSub')))==4
assert imgs(check)==imgs(body)
assert len(check.tables)==len(front.tables)==6
# Data table rows retain the original numbers, rounded only for publication.
expected=[[c.text for r in t.rows[1:] for c in r.cells] for t in body.tables]
actual=[[c.text for r in t.rows[1:] for c in r.cells] for t in check.tables[2:]]
def normalized(vals):
 import re
 return [round(float(v),4) if re.fullmatch(r'\d+\.\d+',v) else v.replace('Source','소스').replace('Target','타깃') for v in vals]
assert [normalized(v) for v in actual]==[normalized(expected[i]) for i in [2,0,1,3]]
for p in [front_path,body_path]:assert hashlib.sha256(p.read_bytes()).hexdigest()==original_hashes[str(p.relative_to(ROOT))]
report={'status':'PASS','supplied_prose_revised_for_review':True,'equations1_and2_repaired':True,'citation_mapping':citation_mapping,'target_data_description':'restored previous wording at user request','cnn_purpose_and_interpretation_added':True,'prospective_cnn_decoder_experiment_mentioned':False,'source_files_unchanged':True,'body_tables_preserved':4,'title_tables_preserved':2,'native_equations_preserved':len(maths(check)),'images_preserved':len(imgs(check)),'author_placeholders_retained':True,'output':str(final),'source_sha256':original_hashes}
(OUT/'final_draft_verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))
print(json.dumps(report,ensure_ascii=False,indent=2))
