"""Render the author's two-page CNN section using the supplied conference template."""
import sys, json, copy
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'vendor/manuscript_deps'))
from docx import Document
from docx.shared import Pt, Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent
RUN = ROOT / 'research/runs/cnn_matched_v1'
doc = Document(ROOT / '논문양식.docx')
sect = copy.deepcopy(doc._element.body.find(qn('w:sectPr')))
for node in list(doc._element.body):
    doc._element.body.remove(node)
doc._element.body.append(sect)
cols = sect.find(qn('w:cols'))
cols.set(qn('w:num'), '2')
cols.set(qn('w:space'), '567')
for h in list(sect):
    if h.tag in [qn('w:headerReference'), qn('w:footerReference')]:
        sect.remove(h)
style = doc.styles['Normal']
style.font.name = 'HY신명조'
style.font.size = Pt(9)
style._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'HY신명조')
style.paragraph_format.line_spacing = 1.2
style.paragraph_format.space_after = Pt(3)

def p(text, heading=False, title=False):
    x = doc.add_paragraph()
    x.paragraph_format.first_line_indent = Pt(9) if not heading else Pt(0)
    if heading:
        x.paragraph_format.keep_with_next = True
        x.paragraph_format.space_before = Pt(7)
    if title:
        x.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = x.add_run(text)
    r.bold = heading
    r.font.size = Pt(12 if title else 10 if heading else 8 if text.startswith(('[1]', '[2]')) else 9)
    return x

def table(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    margins=OxmlElement('w:tblCellMar')
    for edge in ['left','right','top','bottom']:
        item=OxmlElement('w:'+edge);item.set(qn('w:w'),'30');item.set(qn('w:type'),'dxa');margins.append(item)
    t._tbl.tblPr.append(margins)
    widths=([.40,.55,.55,.55,.55,.55] if len(headers)==6 else [.40,1.02,.60,.58,.55] if len(headers)==5 else [3.15/len(headers)]*len(headers))
    for col,width in zip(t.columns,widths):
        col.width = Inches(width)
    borders = OxmlElement('w:tblBorders')
    for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        el = OxmlElement('w:' + edge)
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '4')
        borders.append(el)
    t._tbl.tblPr.append(borders)
    for c, v in zip(t.rows[0].cells, headers):
        c.text = v
    for row in rows:
        for c, v in zip(t.add_row().cells, row):
            c.text = str(v)
    for row in t.rows:
        trPr = row._tr.get_or_add_trPr()
        trPr.append(OxmlElement('w:cantSplit'))
        for c,width in zip(row.cells,widths):
            c.width = Inches(width)
            for x in c.paragraphs:
                x.paragraph_format.space_after = Pt(2)
                x.paragraph_format.space_before = Pt(2)
                for r in x.runs:
                    r.font.size = Pt(8)
    return t
p('III. CNN 기반 조건부 인코더 전이', True, True)
p('3.1 비교 목적 및 문제 정의', True)
p('DINOv2의 전이 범위 비교에 추가하여, CNN에서도 Source 인코더 전이의 효과가 나타나는지 검토한다. 동일한 구조와 Target 입력을 사용하는 무작위 초기화 모델(BASE)과 인코더 전이 모델(ENC)을 비교했다. 입력 표현과 데이터 분할은 DINOv2와 공유하고, 모델별 학습 설정은 구분했다.')
p('모델의 입력은 공간 정보 S, 안테나 패턴 정보 P, 안테나 설정 Q로 구성된다. S는 건물 점유 대리 영상, Tx Gaussian heatmap, Tx log-distance map의 3채널이다. P는 수평 패턴 H, 수직 패턴 V, beamwidth를 영상으로 표현한 n2i의 3채널이다. Q는 height, power, downtilt와 azimuth의 sin·cos 값을 포함한다.')
p('좌표와 패턴 영상은 DINOv2와 같은 전처리 함수로 변환했다. Height, power, downtilt와 H·V는 train set에서 구한 통계로 정규화했다. 정답은 원본 회색조 값을 255로 나누어 0~1 범위로 변환했다.')
p('3.2 조건부 CNN 구조', True)
p('CNN은 채널 수 16·32·64·128의 인코더와 skip connection을 갖는 U-Net[1] 계열 디코더로 구성된다. 각 인코더 블록은 3×3 Conv2d, GroupNorm, SiLU를 두 번 적용하며, 블록 사이에서 max pooling으로 해상도를 줄인다. 디코더는 bilinear interpolation으로 특징을 확대한 뒤 같은 해상도의 인코더 특징과 결합한다. 마지막 1×1 Conv2d는 단일 채널 EMF map을 출력한다.')
p('P는 1×1 Conv2d로 16채널 특징으로 변환하여 첫 인코더 블록의 출력에 더했다. Q는 5→32→256의 MLP로 변환하고, 병목 특징의 128개 채널에 FiLM[2]을 적용했다. 조건화는 다음과 같이 계산한다.')
p('F′ = (1 + γ(Q)) ⊙ F + β(Q).')
p('F는 병목 특징이고 γ와 β는 Q에서 생성한 채널별 배율과 편향이다. 패턴 투영층과 FiLM 출력층은 0으로 초기화했다. 따라서 학습 시작 시 추가 조건화는 공간 특징을 바꾸지 않으며 이후 학습 과정에서 갱신된다. 출력에는 sigmoid와 clamp를 적용하지 않았다.')
p('3.3 인코더 전이와 대응 초기화', True)
p('Source는 일반 영상 사전학습 없이 무작위 초기화한 CNN 전체를 학습했다. Target은 무작위 초기화하는 경우와 Source 인코더만 전이하는 경우를 비교했다. ENC도 디코더와 출력층은 새로 초기화했으며, 인코더를 포함한 전체 파라미터를 학습했다.')
p('Source와 Target 시드는 각각 142·42로 설정했다. BASE와 ENC의 디코더, 출력층, 패턴 투영층, FiLM은 동일하게 초기화했다. ENC는 Source 인코더만 복사한 뒤 전체 모델을 학습했다.')
doc.add_picture(str(OUT / 'cnn_architecture.png'), width=Inches(3.1))
p('그림 1. 공간·패턴·수치 조건을 결합한 CNN 구조. Source 인코더만 전이하고 Target 전체를 학습했다.')
doc.add_page_break()
p('IV. CNN 실험 설정 및 결과', True, True)
p('4.1 공통 데이터와 학습 조건', True)
p('Target은 건물 영상 그룹을 train·validation·test 68·9·9개로 분할했고 표본 수는 609·78·81개이다. Source는 지도 700개를 550·50·100개로 분할하고 지도별 Tx 80개를 사용했다. 두 모델 계열에 같은 데이터 목록을 적용했으며 Target train set 전체로 학습했다.')
p('손실은 점유 대리 영역과 결측 위치를 제외한 유효 픽셀에서 계산했다. CNN에는 L1 손실만 사용했으며, 유효 픽셀 수로 나누어 정규화했다.')
p('L = Σᵢ Mᵢ |ŷᵢ − yᵢ| / Σᵢ Mᵢ.')
p('표 1. CNN 재실험의 학습 설정', True)
table(['항목', '설정'], [['최적화 / 정밀도', 'AdamW / FP32'], ['학습률 / 최저 학습률', '10⁻⁴ / 10⁻⁵ (cosine)'], ['배치 / 누적 / 가중치 감쇠', '8 / 1 / 0'], ['기울기 클리핑', '최대 노름 5'], ['소스 상한 / 워밍업 / patience', '50 / 5 / 15 epoch'], ['타깃 상한 / 워밍업 / patience', '200 / 10 / 50 epoch'], ['조기 종료 최소 epoch', '소스 10 / 타깃 30'], ['증강 / EMA / 출력 클리핑', '모두 미사용']])
p('Validation MAE가 가장 낮은 모델을 선택하여 test set에서 성능을 측정했다. MAE·RMSE·PSNR은 전체 유효 픽셀에서 계산하고 SSIM은 이미지별 값을 평균했다.')
p('4.2 평가 지표와 반복 실험 결과', True)
afile = RUN / 'analysis.json'
if afile.exists():
    assert json.loads((RUN / 'status.json').read_text())['status'] == 'experiments_complete'
    assert json.loads((RUN / 'submission_audit.json').read_text())['status'] == 'PASS'
    a = json.loads(afile.read_text())
    rows = a['rows']
    b, e = rows
    p('표 2. CNN 성능 비교 (Target 시드 42)', True)
    seed42={mode:json.loads((RUN/f'{mode}_seed42/test_metrics.json').read_text())['metrics'] for mode in ['base','encoder']}
    val42={mode:json.loads((RUN/f'{mode}_seed42/summary.json').read_text())['best_validation_mae'] for mode in ['base','encoder']}
    table(['조건','Val MAE','Test MAE','RMSE','SSIM','PSNR'],[[('BASE' if mode=='base' else 'ENC'),f"{val42[mode]:.5f}",f"{seed42[mode]['mae']:.5f}",f"{seed42[mode]['rmse']:.5f}",f"{seed42[mode]['ssim']:.4f}",f"{seed42[mode]['psnr']:.2f}"] for mode in ['base','encoder']])
    x,y=seed42['base'],seed42['encoder']
    rate42=100*(y['mae']-x['mae'])/x['mae']
    p(f"DINOv2와 동일한 Target 시드 42에서 BASE와 ENC의 test MAE는 {x['mae']:.6f}, {y['mae']:.6f}였다. 인코더 전이 후 MAE는 {abs(rate42):.2f}% {'증가' if rate42>=0 else '감소'}했다. RMSE와 PSNR은 BASE가 좋았고 SSIM은 ENC가 좋았다.")

else:
    p('표 2. 최종 테스트 결과 기입 위치', True)
    table(['조건', 'MAE ± SD', 'RMSE', 'SSIM'], [['BASE', '평가 대기', '평가 대기', '평가 대기'], ['ENC', '평가 대기', '평가 대기', '평가 대기']])
    p('마지막 타깃 실행을 완료한 뒤 6개 선택 체크포인트를 고정하여 테스트 성능을 산출한다. 최종 표에는 시드 3개의 MAE 평균과 표준편차, RMSE 및 SSIM 평균을 기입한다. 현재 학습 중인 결과를 완료 수치로 대체하거나 검증 오차를 테스트 오차로 보고하지 않는다.')
    p('추가로 테스트 그룹 9개에서 시드 평균의 그룹별 MAE 차이(BASE−ENC)를 계산하고 10,000회 쌍체 부트스트랩으로 구간을 산출한다. 픽셀 전체 MAE와 그룹 매크로 차이는 가중치가 달라 구별하여 해석한다. 이 분석은 고정된 분할과 소스 체크포인트에 조건부이다.')
p('4.3 DINOv2 비교의 해석 범위', True)
p('DINOv2의 L1→L1 조건과 CNN 모두 인코더 전이만으로 test MAE가 개선되지 않았다. DINOv2에서는 인코더·디코더 전이와 LoRA를 결합한 조건에서 MAE가 낮아졌다. 따라서 이번 결과에서는 Source 학습의 적용 자체보다 전이 범위와 적응 방식을 함께 고려할 필요가 있었다. CNN의 디코더 전이는 비교하지 않았다.')
p('DINOv2는 공개 사전학습 가중치와 LoRA를 사용하고 CNN은 무작위 초기화와 전체 미세조정을 사용했다. 따라서 두 모델의 결과는 초기화와 적응 방식의 차이를 포함하며, 각 모델 내부의 전이 전후 성능을 비교했다.')
p('참고문헌', True)
p('[1] O. Ronneberger, P. Fischer, and T. Brox, “U-Net: Convolutional Networks for Biomedical Image Segmentation,” arXiv:1505.04597, 2015. https://arxiv.org/abs/1505.04597')
p('[2] E. Perez, F. Strub, H. de Vries, V. Dumoulin, and A. Courville, “FiLM: Visual Reasoning with a General Conditioning Layer,” AAAI, 2018. https://arxiv.org/abs/1709.07871')
doc.core_properties.title = 'CNN 기반 조건부 인코더 전이: 담당 본론 초안'
doc.core_properties.author = ''
doc.core_properties.subject = '동일 DINOv2 조건의 CNN 재실험'
path = OUT / 'CNN_담당본론_초안.docx'
doc.save(path)
print(path)
